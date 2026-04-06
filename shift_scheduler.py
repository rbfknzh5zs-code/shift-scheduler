#!/usr/bin/env python3
"""小売店シフト自動作成ツール - PuLPによる最適化"""

import argparse
import calendar
import csv
import dataclasses
import os
import sys
from collections import defaultdict
from datetime import date

from pulp import (
    LpMinimize,
    LpProblem,
    LpVariable,
    LpBinary,
    LpInteger,
    LpStatus,
    lpSum,
    value,
    PULP_CBC_CMD,
)

DAY_NAMES = "月火水木金土日"


@dataclasses.dataclass
class Staff:
    name: str
    max_days_per_week: int
    max_hours_per_week: float
    available_days: set
    patterns: list  # [(start, end), ...]


@dataclasses.dataclass
class CoverageReq:
    day_type: str
    time_start: float
    time_end: float
    min_staff: int


@dataclasses.dataclass
class DayOffRequest:
    name: str
    date: date
    priority: str


@dataclasses.dataclass
class DailyAvailability:
    name: str
    date: date
    time_start: float
    time_end: float


def parse_time_value(text):
    s = str(text).strip()
    if ":" in s:
        hh, mm = s.split(":", 1)
        return int(hh) + int(mm) / 60.0
    return float(s)


def load_staff(path):
    staff_list = []
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            patterns = []
            for p in row["patterns"].split("|"):
                s, e = p.strip().split("-", 1)
                patterns.append((float(s), float(e)))
            staff_list.append(Staff(
                name=row["name"],
                max_days_per_week=int(row["max_days_per_week"]),
                max_hours_per_week=float(row["max_hours_per_week"]),
                available_days=set(row["available_days"]),
                patterns=patterns,
            ))
    return staff_list


def load_coverage(path):
    reqs = []
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            reqs.append(CoverageReq(
                day_type=row["day_type"],
                time_start=parse_time_value(row["time_start"]),
                time_end=parse_time_value(row["time_end"]),
                min_staff=int(row["min_staff"]),
            ))
    return reqs


def load_requests(path):
    reqs = []
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            y, m, d = row["date"].split("-")
            reqs.append(DayOffRequest(
                name=row["name"],
                date=date(int(y), int(m), int(d)),
                priority=row["priority"],
            ))
    return reqs


def load_daily_availability(path):
    rows = []
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            if not row.get("name") or not row.get("date"):
                continue
            y, m, d = row["date"].split("-")
            ts = row.get("time_start", "").strip()
            te = row.get("time_end", "").strip()
            if not ts or not te:
                continue
            rows.append(DailyAvailability(
                name=row["name"].strip(),
                date=date(int(y), int(m), int(d)),
                time_start=parse_time_value(ts),
                time_end=parse_time_value(te),
            ))
    return rows


def get_day_type(d, holidays):
    if d in holidays or d.weekday() == 6:
        return "日祝"
    if d.weekday() == 5:
        return "土曜"
    return "平日"


def day_name(d):
    return DAY_NAMES[d.weekday()]


def get_month_dates(year, month):
    _, last = calendar.monthrange(year, month)
    return [date(year, month, day) for day in range(1, last + 1)]


def get_iso_weeks(dates):
    weeks = defaultdict(list)
    for d in dates:
        weeks[d.isocalendar()[1]].append(d)
    return dict(weeks)


def fmt_time(pattern):
    if pattern is None:
        return "休み"
    return f"{pattern[0]:g}-{pattern[1]:g}"


def pattern_hours(pattern):
    if pattern is None:
        return 0.0
    return pattern[1] - pattern[0]


def resolve_target_year_month(year, month):
    if year is not None and month is not None:
        return year, month
    if year is None and month is None:
        today = date.today()
        if today.month == 12:
            return today.year + 1, 1
        return today.year, today.month + 1
    raise ValueError("--year と --month はセットで指定してください。省略時は次月を自動対象にします。")


def parse_holidays(holidays_arg):
    holidays = set()
    if not holidays_arg:
        return holidays
    for h in holidays_arg.split(","):
        h = h.strip()
        if not h:
            continue
        y, m, d = h.split("-")
        holidays.add(date(int(y), int(m), int(d)))
    return holidays


def get_jp_holidays_for_dates(dates):
    try:
        import jpholiday
    except ImportError:
        return set(), False
    return {d for d in dates if jpholiday.is_holiday(d)}, True


def solve_schedule(
    year, month, staff_list, coverage_reqs, day_off_requests, holidays,
    daily_availability=None,
    max_consecutive=5, w_request=100, w_fairness=10, w_consecutive=10, timeout=60,
):
    dates = get_month_dates(year, month)
    date_set = set(dates)
    weeks = get_iso_weeks(dates)

    # カバレッジマップ: (day_type, hour) -> min_staff
    coverage_map = {}
    for req in coverage_reqs:
        for h in range(int(req.time_start), int(req.time_end)):
            key = (req.day_type, h)
            coverage_map[key] = max(coverage_map.get(key, 0), req.min_staff)
    all_hours = sorted(set(h for (_, h) in coverage_map))

    must_off = defaultdict(set)
    wish_off = defaultdict(set)
    for r in day_off_requests:
        if r.date in date_set:
            if r.priority == "必須":
                must_off[r.name].add(r.date)
            else:
                wish_off[r.name].add(r.date)

    available_window = {}
    for row in (daily_availability or []):
        if row.date in date_set:
            available_window[(row.name, row.date)] = (row.time_start, row.time_end)

    prob = LpProblem("ShiftScheduler", LpMinimize)

    # 決定変数: x[si, d_idx, pi] スタッフsiが日d_idxにパターンpiで勤務
    x = {}
    for si, s in enumerate(staff_list):
        for di, d in enumerate(dates):
            for pi in range(len(s.patterns)):
                x[(si, di, pi)] = LpVariable(f"x_{si}_{di}_{pi}", cat=LpBinary)

    # 補助変数: work[si, d_idx] 出勤するか
    work = {}
    for si, s in enumerate(staff_list):
        for di in range(len(dates)):
            work[(si, di)] = LpVariable(f"w_{si}_{di}", cat=LpBinary)

    # work = sum(x)
    for si, s in enumerate(staff_list):
        for di in range(len(dates)):
            prob += work[(si, di)] == lpSum(
                x[(si, di, pi)] for pi in range(len(s.patterns))
            )

    # --- ハード制約 ---

    # C1: 1日1パターンまで
    for si, s in enumerate(staff_list):
        for di in range(len(dates)):
            if len(s.patterns) > 1:
                prob += lpSum(x[(si, di, pi)] for pi in range(len(s.patterns))) <= 1

    # C2: カバレッジ (各時間帯の必要人数)
    for di, d in enumerate(dates):
        dt = get_day_type(d, holidays)
        for h in all_hours:
            min_staff = coverage_map.get((dt, h), 0)
            if min_staff <= 0:
                continue
            covering = []
            for si, s in enumerate(staff_list):
                for pi, pat in enumerate(s.patterns):
                    if pat[0] <= h and pat[1] > h:
                        covering.append(x[(si, di, pi)])
            # 候補が0人でも制約は必ず追加し、要求を満たせないケースを
            # "不足を見逃す" のではなく "Infeasible" として扱う。
            prob += lpSum(covering) >= min_staff, f"cov_{di}_{h}"

    # C3: 勤務可能曜日
    for si, s in enumerate(staff_list):
        for di, d in enumerate(dates):
            if day_name(d) not in s.available_days:
                prob += work[(si, di)] == 0

    # C4: 週あたり勤務日数上限
    for si, s in enumerate(staff_list):
        for wn, wdates in weeks.items():
            prob += lpSum(
                work[(si, dates.index(d))] for d in wdates
            ) <= s.max_days_per_week, f"maxd_{si}_{wn}"

    # C5: 週あたり勤務時間上限
    for si, s in enumerate(staff_list):
        for wn, wdates in weeks.items():
            prob += lpSum(
                (pat[1] - pat[0]) * x[(si, dates.index(d), pi)]
                for d in wdates
                for pi, pat in enumerate(s.patterns)
            ) <= s.max_hours_per_week, f"maxh_{si}_{wn}"

    # C6: 必須休
    for si, s in enumerate(staff_list):
        for d in must_off.get(s.name, set()):
            if d in date_set:
                di = dates.index(d)
                prob += work[(si, di)] == 0

    # C7: 日別の勤務可能時間
    for si, s in enumerate(staff_list):
        for di, d in enumerate(dates):
            window = available_window.get((s.name, d))
            if window is None:
                continue
            ws, we = window
            if ws >= we:
                prob += work[(si, di)] == 0
                continue
            for pi, pat in enumerate(s.patterns):
                if pat[0] < ws or pat[1] > we:
                    prob += x[(si, di, pi)] == 0

    # --- ソフト制約 ---

    # 希望休違反
    wish_penalty = []
    for si, s in enumerate(staff_list):
        for d in wish_off.get(s.name, set()):
            if d in date_set:
                wish_penalty.append(work[(si, dates.index(d))])

    # 公平性: 時間ベースの偏差
    total_hours_var = {}
    for si, s in enumerate(staff_list):
        total_hours_var[si] = lpSum(
            (pat[1] - pat[0]) * x[(si, di, pi)]
            for di in range(len(dates))
            for pi, pat in enumerate(s.patterns)
        )

    total_coverage_hours = sum(
        coverage_map.get((get_day_type(d, holidays), h), 0)
        for d in dates for h in all_hours
    )
    total_capacity = sum(s.max_hours_per_week for s in staff_list)

    deviation = {}
    for si, s in enumerate(staff_list):
        target = total_coverage_hours * (s.max_hours_per_week / total_capacity) if total_capacity > 0 else 0
        deviation[si] = LpVariable(f"dev_{si}", lowBound=0)
        prob += deviation[si] >= total_hours_var[si] - target
        prob += deviation[si] >= target - total_hours_var[si]

    # 連勤ペナルティ
    consec_vars = []
    for si in range(len(staff_list)):
        for i in range(len(dates) - max_consecutive):
            cv = LpVariable(f"con_{si}_{i}", lowBound=0, cat=LpInteger)
            prob += cv >= lpSum(work[(si, di)] for di in range(i, i + max_consecutive + 1)) - max_consecutive
            consec_vars.append(cv)

    prob += (
        w_request * lpSum(wish_penalty)
        + w_fairness * lpSum(deviation[si] for si in range(len(staff_list)))
        + w_consecutive * lpSum(consec_vars)
    )

    prob.solve(PULP_CBC_CMD(msg=0, timeLimit=timeout))

    if prob.status != 1:
        return None, LpStatus[prob.status]

    # 結果抽出: schedule[(name, date)] = (start, end) or None
    schedule = {}
    for si, s in enumerate(staff_list):
        for di, d in enumerate(dates):
            assigned = None
            for pi, pat in enumerate(s.patterns):
                if value(x[(si, di, pi)]) > 0.5:
                    assigned = pat
                    break
            schedule[(s.name, d)] = assigned

    return schedule, "Optimal"


def write_csv(schedule, year, month, staff_list, day_off_requests, output_path):
    dates = get_month_dates(year, month)
    names = [s.name for s in staff_list]

    req_map = defaultdict(lambda: defaultdict(str))
    for r in day_off_requests:
        if r.date.year == year and r.date.month == month:
            req_map[r.name][r.date] = r.priority

    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        # ヘッダ: 日付行
        header_date = [""] + [d.strftime("%m/%d") for d in dates] + ["合計時間"]
        header_day = [""] + [day_name(d) for d in dates] + [""]
        writer.writerow(header_date)
        writer.writerow(header_day)

        # スタッフ行
        for name in names:
            row = [name]
            total_h = 0.0
            for d in dates:
                pat = schedule.get((name, d))
                row.append(fmt_time(pat))
                total_h += pattern_hours(pat)
            row.append(f"{total_h:g}")
            writer.writerow(row)

        # 店舗人時
        writer.writerow([])
        ph_row = ["店舗人時"]
        grand_total = 0.0
        for d in dates:
            daily = sum(pattern_hours(schedule.get((n, d))) for n in names)
            ph_row.append(f"{daily:g}")
            grand_total += daily
        ph_row.append(f"{grand_total:g}")
        writer.writerow(ph_row)

        # 希望休反映
        req_row = ["希望休反映"]
        for d in dates:
            req_row.append("")
        req_row.append("")
        writer.writerow([])
        for name in names:
            total_req = 0
            fulfilled = 0
            for d_req, pri in req_map[name].items():
                total_req += 1
                if schedule.get((name, d_req)) is None:
                    fulfilled += 1
            if total_req > 0:
                writer.writerow([name, f"{fulfilled}/{total_req}反映"])
            else:
                writer.writerow([name, "希望なし"])


def write_xlsx(schedule, year, month, staff_list, day_off_requests, holidays, coverage_reqs, output_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    dates = get_month_dates(year, month)
    names = [s.name for s in staff_list]

    req_map = defaultdict(lambda: defaultdict(str))
    for r in day_off_requests:
        if r.date.year == year and r.date.month == month:
            req_map[r.name][r.date] = r.priority

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}年{month}月シフト表"

    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    sat_font = Font(color="0000FF", size=10)
    sun_font = Font(color="FF0000", size=10)
    bold_font = Font(bold=True, size=10)
    normal_font = Font(size=10)
    weekend_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    off_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    morning_fill = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
    evening_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    total_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")

    num_dates = len(dates)
    total_col = num_dates + 2

    # 行1: 日付ヘッダ
    ws.cell(row=1, column=1, value="スタッフ").font = hdr_font
    ws.cell(row=1, column=1).fill = hdr_fill
    ws.cell(row=1, column=1).alignment = center
    ws.cell(row=1, column=1).border = thin

    for ci, d in enumerate(dates, 2):
        cell = ws.cell(row=1, column=ci, value=d.strftime("%m/%d"))
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center
        cell.border = thin

    cell = ws.cell(row=1, column=total_col, value="合計時間")
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = center
    cell.border = thin

    # 行2: 曜日ヘッダ
    ws.cell(row=2, column=1, value="").border = thin
    for ci, d in enumerate(dates, 2):
        dn = day_name(d)
        cell = ws.cell(row=2, column=ci, value=dn)
        cell.alignment = center
        cell.border = thin
        if d.weekday() == 5:
            cell.font = sat_font
            cell.fill = weekend_fill
        elif d.weekday() == 6:
            cell.font = sun_font
            cell.fill = weekend_fill
        else:
            cell.font = normal_font
    ws.cell(row=2, column=total_col, value="").border = thin

    # スタッフ行
    for ri, name in enumerate(names, 3):
        cell = ws.cell(row=ri, column=1, value=name)
        cell.font = bold_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin

        total_h = 0.0
        for ci, d in enumerate(dates, 2):
            pat = schedule.get((name, d))
            val = fmt_time(pat)
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.alignment = center
            cell.border = thin
            cell.font = normal_font

            if pat is None:
                cell.fill = off_fill
            elif pat[0] < 12:
                cell.fill = morning_fill
            else:
                cell.fill = evening_fill

            total_h += pattern_hours(pat)

        cell = ws.cell(row=ri, column=total_col, value=round(total_h, 2))
        cell.alignment = center
        cell.border = thin
        cell.font = bold_font
        cell.fill = total_fill

    # 店舗人時行
    ph_row = len(names) + 4
    cell = ws.cell(row=ph_row, column=1, value="店舗人時")
    cell.font = bold_font
    cell.border = thin

    grand_total = 0.0
    for ci, d in enumerate(dates, 2):
        daily = sum(pattern_hours(schedule.get((n, d))) for n in names)
        cell = ws.cell(row=ph_row, column=ci, value=round(daily, 2))
        cell.alignment = center
        cell.border = thin
        cell.font = bold_font
        grand_total += daily

    cell = ws.cell(row=ph_row, column=total_col, value=round(grand_total, 2))
    cell.alignment = center
    cell.border = thin
    cell.font = Font(bold=True, color="FF0000", size=11)
    cell.fill = total_fill

    # 列幅調整
    ws.column_dimensions["A"].width = 14
    for ci in range(2, total_col + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 12

    # 行高さ
    for ri in range(1, ph_row + 1):
        ws.row_dimensions[ri].height = 22

    wb.save(output_path)


def print_summary(schedule, year, month, staff_list, coverage_reqs, day_off_requests, holidays):
    dates = get_month_dates(year, month)
    names = [s.name for s in staff_list]
    staff_map = {s.name: s for s in staff_list}

    req_map = defaultdict(lambda: defaultdict(str))
    for r in day_off_requests:
        if r.date.year == year and r.date.month == month:
            req_map[r.name][r.date] = r.priority

    coverage_map = {}
    for req in coverage_reqs:
        for h in range(int(req.time_start), int(req.time_end)):
            key = (req.day_type, h)
            coverage_map[key] = max(coverage_map.get(key, 0), req.min_staff)
    all_hours = sorted(set(h for (_, h) in coverage_map))

    print(f"\n=== シフト作成結果 ({year}年{month}月) ===\n")

    print("【スタッフ別集計】")
    for name in names:
        s = staff_map[name]
        work_days = sum(1 for d in dates if schedule.get((name, d)) is not None)
        total_h = sum(pattern_hours(schedule.get((name, d))) for d in dates)
        total_req = 0
        fulfilled = 0
        for d_req, pri in req_map[name].items():
            total_req += 1
            if schedule.get((name, d_req)) is None:
                fulfilled += 1
        req_str = f"{fulfilled}/{total_req}反映" if total_req > 0 else "希望なし"
        print(f"  {name}: {work_days}日 {total_h:g}h (上限: 週{s.max_days_per_week}日/{s.max_hours_per_week:g}h) 希望休: {req_str}")

    print("\n【カバレッジ不足】")
    shortage = False
    for d in dates:
        dt = get_day_type(d, holidays)
        for h in all_hours:
            need = coverage_map.get((dt, h), 0)
            if need <= 0:
                continue
            actual = 0
            for name in names:
                pat = schedule.get((name, d))
                if pat and pat[0] <= h and pat[1] > h:
                    actual += 1
            if actual < need:
                print(f"  {d}({day_name(d)}) {h}時台: {actual}/{need}人")
                shortage = True
    if not shortage:
        print("  なし")

    print("\n【連勤状況】")
    for name in names:
        max_c = current = 0
        for d in dates:
            if schedule.get((name, d)) is not None:
                current += 1
                max_c = max(max_c, current)
            else:
                current = 0
        if max_c >= 4:
            print(f"  {name}: 最大{max_c}連勤")

    grand = sum(pattern_hours(schedule.get((n, d))) for n in names for d in dates)
    print(f"\n店舗人時合計: {grand:g}h\n")


def main():
    parser = argparse.ArgumentParser(description="小売店シフト自動作成ツール")
    parser.add_argument("--year", type=int, default=None)
    parser.add_argument("--month", type=int, default=None)
    parser.add_argument("--staff", default="sample_data/staff.csv")
    parser.add_argument("--coverage", default="sample_data/coverage.csv")
    parser.add_argument("--requests", default="sample_data/requests.csv")
    parser.add_argument("--daily-availability", default=None)
    parser.add_argument("--output", default=None)
    parser.add_argument("--max-consecutive", type=int, default=5)
    parser.add_argument("--holidays", default="")
    parser.add_argument("--auto-jp-holidays", action=argparse.BooleanOptionalAction, default=True)
    parser.add_argument("--format", choices=["csv", "xlsx", "both"], default="csv")
    parser.add_argument("--timeout", type=int, default=60)
    args = parser.parse_args()

    try:
        target_year, target_month = resolve_target_year_month(args.year, args.month)
    except ValueError as e:
        print(f"引数エラー: {e}")
        sys.exit(2)

    target_dates = get_month_dates(target_year, target_month)
    manual_holidays = parse_holidays(args.holidays)
    holidays = set(manual_holidays)
    auto_holidays = set()
    auto_holiday_available = True
    if args.auto_jp_holidays:
        auto_holidays, auto_holiday_available = get_jp_holidays_for_dates(target_dates)
        holidays |= auto_holidays

    output_csv = None
    output_xlsx = None
    if args.format == "both":
        if args.output is None:
            output_csv = f"shift_output_{target_year}{target_month:02d}.csv"
            output_xlsx = f"shift_output_{target_year}{target_month:02d}.xlsx"
        else:
            base, ext = os.path.splitext(args.output)
            if ext.lower() in {".csv", ".xlsx"}:
                out_base = base
            else:
                out_base = args.output
            output_csv = f"{out_base}.csv"
            output_xlsx = f"{out_base}.xlsx"
    else:
        if args.output is None:
            args.output = f"shift_output_{target_year}{target_month:02d}.{args.format}"
        if args.format == "csv":
            output_csv = args.output
        else:
            output_xlsx = args.output

    print("データ読み込み中...")
    staff_list = load_staff(args.staff)
    coverage_reqs = load_coverage(args.coverage)
    requests_data = load_requests(args.requests)
    availability_data = load_daily_availability(args.daily_availability) if args.daily_availability else []

    print(f"スタッフ: {len(staff_list)}人")
    print(f"カバレッジ要件: {len(coverage_reqs)}件")
    print(f"希望休: {len(requests_data)}件")
    print(f"日別可能時間: {len(availability_data)}件")
    print(f"対象: {target_year}年{target_month}月")
    if args.auto_jp_holidays:
        if auto_holiday_available:
            print(f"祝日自動判定: {len(auto_holidays)}日")
        else:
            print("祝日自動判定: jpholiday 未インストールのため無効")
    print(f"手動祝日指定: {len(manual_holidays)}日")
    print(f"休日扱い合計: {len(holidays)}日\n")

    print("最適化計算中...")
    schedule, status = solve_schedule(
        year=target_year,
        month=target_month,
        staff_list=staff_list,
        coverage_reqs=coverage_reqs,
        day_off_requests=requests_data,
        holidays=holidays,
        daily_availability=availability_data,
        max_consecutive=args.max_consecutive,
        timeout=args.timeout,
    )

    if schedule is None:
        print(f"\nエラー: シフト作成不可 (ステータス: {status})")
        print("スタッフの人数や勤務条件を見直してください。")
        sys.exit(1)

    print(f"ステータス: {status}")

    if output_xlsx:
        write_xlsx(schedule, target_year, target_month, staff_list, requests_data, holidays, coverage_reqs, output_xlsx)
    if output_csv:
        write_csv(schedule, target_year, target_month, staff_list, requests_data, output_csv)

    if args.format == "both":
        print(f"出力: {output_csv}, {output_xlsx}")
    elif output_csv:
        print(f"出力: {output_csv}")
    else:
        print(f"出力: {output_xlsx}")
    print_summary(schedule, target_year, target_month, staff_list, coverage_reqs, requests_data, holidays)


if __name__ == "__main__":
    main()
